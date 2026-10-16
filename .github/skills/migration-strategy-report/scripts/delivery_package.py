"""Render and verify an optional execution summary from validated report data."""

import argparse
import json
from pathlib import Path
import sys
from zipfile import ZipFile

from report_data import PILLARS, load_model, require, validate_model


TITLES = ("Execution Summary", "Microsoft Scope", "Partner and Unresolved Scope",
          "Top Five Waves and Blockers")
SHEETS = ("Open Questions", "Delivery Blockers", "Scope and Waves", "Validation Steps")


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
                    encoding="utf-8")


def cell_text(value):
    # All customer strings are literal text, never spreadsheet formulas.
    result = str(value)
    require(len(result) <= 32767, "Workbook cell exceeds Excel's text limit; split the evidence")
    return result


def add_row(sheet, values):
    sheet.append([cell_text(value) for value in values])
    for cell in sheet[sheet.max_row]:
        cell.data_type = "s"


def scope_lines(summary, owners):
    lines = []
    for pillar in PILLARS:
        values = summary[pillar]
        if values["portfolio"]:
            counts = values["ownership"]
            split = ", ".join("{}: {}".format(owner, counts[owner])
                              for owner in owners if counts.get(owner))
            lines.append("{}: {} ({} in scope)".format(
                pillar.capitalize(), split or "None", values["in_scope"]))
    return lines or ["No structured inventory supplied; see assessment limitations."]


def render_package(model, output):
    summary = validate_model(model)
    require(bool(model["records"]), "Execution package requires structured inventory records")
    # Import before creating output so a missing optional dependency leaves no package.
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font

    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    steps = []
    validation = {"outcome": "FAILED", "validation_steps": steps}
    validation_path = output / "Validation.json"
    write_json(validation_path, validation)

    def check(identifier, method, passed):
        steps.append({"check": identifier, "method": method,
                      "status": "PASS" if passed else "FAIL"})
        require(passed, "Delivery validation failed: {}".format(identifier))

    try:
        conditional = bool(model.get("questions") or model.get("blockers")) or any(
            row["in_scope"] and (row["ownership"] == "Unknown" or row["confidence"] == "Low")
            for row in model["records"]
        ) or any(values["unassigned_migrations"] for values in summary.values())
        outcome = "CONDITIONAL" if conditional else "READY"
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        wave_lines = []
        for wave in model.get("waves", [])[:5]:
            counts = {pillar: sum(ref["pillar"] == pillar for ref in wave["record_refs"])
                      for pillar in PILLARS}
            wave_lines.append("{}: {} ({})".format(
                wave["id"], wave["name"],
                "; ".join("{} {}".format(count, pillar) for pillar, count in counts.items() if count)
            ))
        blockers = model.get("blockers", [])
        wave_lines += ["Blocker {}: {}".format(item["id"], item["text"]) for item in blockers[:5]]
        wave_lines.append("Full inventory, waves and blockers: Execution_Workbook.xlsx")
        bodies = (
            [model["customer"], "Package status: {}".format(outcome),
             "Inventory units remain separate: applications / DB instances / VMs.",
             "Recommendations require customer approval before execution.",
             "Full evidence: Execution_Workbook.xlsx and Traceability.json"],
            scope_lines(summary, ("Microsoft",)),
            scope_lines(summary, ("Partner", "Unknown", "No Migration Needed")),
            wave_lines,
        )
        for title, lines in zip(TITLES, bodies):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text_frame.clear()
            frame = slide.shapes.add_textbox(Inches(0.6), Inches(1.5),
                                            Inches(12.1), Inches(5.4)).text_frame
            frame.word_wrap = True
            # Keep the slide bounded; exact unabridged text remains in the workbook/JSON.
            frame.text = "\n".join(line if len(line) <= 110 else line[:107] + "..." for line in lines)
            for paragraph in frame.paragraphs:
                paragraph.font.size = Pt(16)
                paragraph.line_spacing = 1.0
                paragraph.space_after = Pt(2)
        deck_path = output / "Execution_Summary.pptx"
        prs.save(deck_path)

        wb = Workbook()
        wb.remove(wb.active)
        for name in SHEETS:
            wb.create_sheet(name)
        for name, field in ((SHEETS[0], "questions"), (SHEETS[1], "blockers")):
            sheet = wb[name]
            add_row(sheet, ["ID", "Question / Blocker", "Accountable Owner", "Evidence"])
            for item in model.get(field, []):
                add_row(sheet, [item["id"], item["text"], item["owner"], "\n".join(item["evidence"])])
        scope = wb[SHEETS[2]]
        add_row(scope, ["Pillar", "ID", "Name", "In Scope", "Ownership", "Strategy",
                        "Target (Proposed)", "Confidence", "Rule", "Rationale", "Wave", "Evidence"])
        assignments = {(ref["pillar"], ref["id"]): wave["id"]
                       for wave in model.get("waves", []) for ref in wave["record_refs"]}
        for row in model["records"]:
            wave = assignments.get((row["pillar"], row["id"]),
                                   "Unassigned" if row["in_scope"] and row["migration_required"] else "N/A")
            add_row(scope, [row["pillar"], row["id"], row["name"], row["in_scope"], row["ownership"],
                            row.get("strategy", ""), row.get("target", "Undecided"), row["confidence"],
                            row["rule"], row["rationale"], wave, "\n".join(row["evidence"])])
        for sheet in wb:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 28
            for row in sheet:
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

        trace = {"model": model, "summary": summary}
        trace_path = output / "Traceability.json"
        write_json(trace_path, trace)
        check("traceability-round-trip", "Parse saved JSON and compare model and computed summaries",
              json.loads(trace_path.read_text(encoding="utf-8")) == trace)
        with ZipFile(deck_path) as package:
            check("pptx-zip-integrity", "ZIP CRC and required presentation part",
                  package.testzip() is None and "ppt/presentation.xml" in package.namelist())
        reopened = Presentation(deck_path)
        check("pptx-slide-count", "Reopen saved PPTX with python-pptx", len(reopened.slides) == 4)
        for number, title in enumerate(TITLES):
            expected_body = "\n".join(line if len(line) <= 110 else line[:107] + "..."
                                      for line in bodies[number])
            check("pptx-slide-{}".format(number + 1), "Verify saved slide title and computed content",
                  reopened.slides[number].shapes.title.text == title
                  and reopened.slides[number].shapes[-1].text == expected_body)
        workbook_path = output / "Execution_Workbook.xlsx"

        def save_workbook():
            sheet = wb[SHEETS[3]]
            if sheet.max_row:
                sheet.delete_rows(1, sheet.max_row)
            add_row(sheet, ["Check", "Status", "Validation Method"])
            for step in steps:
                add_row(sheet, [step["check"], step["status"], step["method"]])
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            wb.save(workbook_path)

        save_workbook()
        with ZipFile(workbook_path) as package:
            check("xlsx-zip-integrity", "ZIP CRC and required workbook part",
                  package.testzip() is None and "xl/workbook.xml" in package.namelist())
        saved = load_workbook(workbook_path)
        try:
            check("xlsx-sheets", "Reopen XLSX and compare required sheet names", saved.sheetnames == list(SHEETS))
            expected_rows = (len(model.get("questions", [])) + 1,
                             len(model.get("blockers", [])) + 1, len(model["records"]) + 1)
            for name, expected in zip(SHEETS[:3], expected_rows):
                expected_values = [tuple("" if value is None else value for value in row)
                                   for row in wb[name].values]
                actual_values = [tuple("" if value is None else value for value in row)
                                 for row in saved[name].values]
                check("xlsx-{}".format(name), "Verify saved worksheet row count and full evidence content",
                      saved[name].max_row == expected and actual_values == expected_values)
            check("xlsx-literal-text", "Verify customer cells are not executable formulas",
                  all(cell.data_type != "f" for sheet in saved for row in sheet for cell in row))
        finally:
            saved.close()
        save_workbook()
        saved = load_workbook(workbook_path, read_only=True)
        try:
            expected = [("Check", "Status", "Validation Method")] + [
                (step["check"], step["status"], step["method"]) for step in steps]
            require(list(saved[SHEETS[3]].values) == expected,
                    "Saved workbook validation evidence differs from executed checks")
        finally:
            saved.close()
        validation["outcome"] = outcome
    except (OSError, ValueError) as error:
        validation["error"] = str(error)
        raise
    finally:
        write_json(validation_path, validation)
    return validation


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    try:
        result = render_package(load_model(args.input), args.output)
    except (OSError, ValueError, ImportError) as error:
        print("Execution package FAILED: {}".format(error), file=sys.stderr)
        return 1
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
