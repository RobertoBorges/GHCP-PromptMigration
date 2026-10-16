"""Focused regression tests for portfolio reporting evidence and Office delivery."""

from copy import deepcopy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
TOOLS = ROOT / ".github" / "skills" / "migration-strategy-report" / "scripts"
sys.path.insert(0, str(TOOLS))

from report_data import load_model, validate_model
from delivery_package import render_package


def record(pillar="applications", identifier="ONE", **changes):
    row = {
        "pillar": pillar, "id": identifier, "name": "Orders",
        "in_scope": True, "ownership": "Microsoft", "strategy": "Rehost",
        "migration_required": True, "confidence": "High", "rule": "SUPPORTED-PATTERN",
        "rationale": "Verified pattern under the supplied scope reference.",
        "evidence": ["D1: Inventory.xlsx / Apps / row 2"],
    }
    row.update(changes)
    return row


def model(*records):
    return {"schema_version": "1.0", "customer": "Contoso", "language": "en-US",
            "records": list(records), "waves": [], "questions": [], "blockers": []}


def assign_all(data):
    data["waves"] = [{"id": "W1", "name": "Pilot",
                      "record_refs": [{"pillar": row["pillar"], "id": row["id"]}
                                      for row in data["records"]
                                      if row["in_scope"] and row["migration_required"]]}]
    return data


class EvidenceModelTests(unittest.TestCase):
    def test_distinct_units_scope_and_same_id_across_pillars(self):
        data = model(record(), record("databases"), record("infrastructure"),
                     record(identifier="OUT", in_scope=False),
                     record(identifier="SAAS", ownership="No Migration Needed",
                            strategy="Retain", migration_required=False))
        summary = validate_model(data)
        self.assertEqual(summary["applications"]["portfolio"], 3)
        self.assertEqual(summary["applications"]["in_scope"], 2)
        self.assertEqual(summary["applications"]["ownership"], {"Microsoft": 1, "No Migration Needed": 1})
        self.assertEqual(summary["applications"]["strategies"], {"Rehost": 1, "Retain": 1})
        self.assertEqual(summary["databases"]["in_scope"], 1)
        self.assertEqual(summary["infrastructure"]["in_scope"], 1)
        self.assertNotIn("grand_total", summary)

    def test_duplicate_rows_cannot_inflate_totals(self):
        with self.assertRaisesRegex(ValueError, "duplicates a pillar/id"):
            validate_model(model(record(), record()))

    def test_invalid_dispositions_and_missing_evidence_fail(self):
        cases = (
            {"in_scope": "false"}, {"evidence": []}, {"confidence": "Certain"},
            {"ownership": "No Migration Needed"}, {"strategy": "Retire"},
            {"name": "Invalid\x00value"}, {"rule": ""},
        )
        for changes in cases:
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                validate_model(model(record(**changes)))
        with self.assertRaisesRegex(ValueError, "application-only"):
            validate_model(model(record("databases", ownership="No Migration Needed",
                                        migration_required=False)))

    def test_eight_r_mode_requires_explicit_selection(self):
        data = model(record(strategy="Rearchitect"))
        with self.assertRaises(ValueError):
            validate_model(data)
        data["strategy_mode"] = "8R"
        self.assertEqual(validate_model(data)["applications"]["strategies"], {"Rearchitect": 1})

    def test_wave_assignments_require_unique_scoped_migrating_records(self):
        base = assign_all(model(record()))
        self.assertEqual(validate_model(base)["applications"]["unassigned_migrations"], 0)
        for extra in (
            {"id": "W2", "name": "Duplicate", "record_refs": base["waves"][0]["record_refs"]},
            {"id": "W2", "name": "Unknown", "record_refs": [{"pillar": "applications", "id": "MISSING"}]},
        ):
            data = deepcopy(base)
            data["waves"].append(extra)
            with self.assertRaises(ValueError):
                validate_model(data)
        for changes in ({"in_scope": False}, {"strategy": "Retain", "migration_required": False}):
            data = deepcopy(base)
            data["records"][0].update(changes)
            with self.assertRaisesRegex(ValueError, "out-of-scope or non-migrating"):
                validate_model(data)

    def test_host_join_must_resolve_and_does_not_create_servers(self):
        data = model(record(host_ids=["VM1"]), record("infrastructure", "VM1"))
        self.assertEqual(validate_model(data)["infrastructure"]["portfolio"], 1)
        data["records"][0]["host_ids"] = ["MISSING"]
        with self.assertRaisesRegex(ValueError, "unknown VM"):
            validate_model(data)

    def test_duplicate_json_keys_and_invalid_cli_input_fail(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "invalid.json"
            source.write_text('{"records": [], "records": []}', encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Duplicate JSON"):
                load_model(source)
            result = subprocess.run([sys.executable, "-B", str(TOOLS / "report_data.py"), str(source)],
                                    capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("validation FAILED", result.stderr)
            self.assertEqual(result.stdout, "")


class DeliveryTests(unittest.TestCase):
    def test_office_outputs_reconcile_and_preserve_literal_evidence(self):
        from openpyxl import load_workbook
        from pptx import Presentation

        data = assign_all(model(record(name="=HYPERLINK(\"https://invalid.example\")"),
                                record("databases"), record("infrastructure")))
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / "package"
            result = render_package(data, target)
            self.assertEqual(result["outcome"], "READY")
            self.assertTrue(all(step["status"] == "PASS" for step in result["validation_steps"]))
            self.assertEqual(len(Presentation(target / "Execution_Summary.pptx").slides), 4)
            trace = json.loads((target / "Traceability.json").read_text(encoding="utf-8"))
            self.assertEqual(trace["model"], data)
            self.assertEqual(trace["summary"], validate_model(data))
            saved = load_workbook(target / "Execution_Workbook.xlsx")
            try:
                cell = saved["Scope and Waves"]["C2"]
                self.assertEqual(cell.value, data["records"][0]["name"])
                self.assertEqual(cell.data_type, "s")
                self.assertEqual(saved["Scope and Waves"].max_row, 4)
                self.assertEqual(saved["Validation Steps"].max_row, len(result["validation_steps"]) + 1)
            finally:
                saved.close()
            self.assertEqual(json.loads((target / "Validation.json").read_text(encoding="utf-8")), result)
            original = (target / "Traceability.json").read_bytes()
            with self.assertRaises(FileExistsError):
                render_package(data, target)
            self.assertEqual((target / "Traceability.json").read_bytes(), original)

    def test_qualifications_remain_conditional(self):
        cases = (
            model(record()),
            assign_all(model(record(ownership="Unknown"))),
            assign_all(model(record(confidence="Low"))),
        )
        qualified = assign_all(model(record()))
        qualified["blockers"] = [{"id": "B1", "text": "Approval pending", "owner": "Application owner",
                                  "evidence": ["D1 row 2"]}]
        cases += (qualified,)
        with tempfile.TemporaryDirectory() as temp:
            for number, data in enumerate(cases):
                with self.subTest(number=number):
                    self.assertEqual(render_package(data, Path(temp) / str(number))["outcome"],
                                     "CONDITIONAL")

    def test_invalid_model_creates_no_package(self):
        with tempfile.TemporaryDirectory() as temp:
            for number, data in enumerate((model(), model(record(evidence=[])))):
                target = Path(temp) / str(number)
                with self.assertRaises(ValueError):
                    render_package(data, target)
                self.assertFalse(target.exists())

    def test_oversize_workbook_text_fails_explicitly_not_truncated(self):
        data = assign_all(model(record(name="A" * 32768)))
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / "failed"
            with self.assertRaisesRegex(ValueError, "text limit"):
                render_package(data, target)
            result = json.loads((target / "Validation.json").read_text(encoding="utf-8"))
            self.assertEqual(result["outcome"], "FAILED")
            self.assertIn("text limit", result["error"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
